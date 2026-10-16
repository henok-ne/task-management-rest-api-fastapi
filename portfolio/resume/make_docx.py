#!/usr/bin/env python3
"""
Builds Henok Mesfin's resume as a Word (.docx) file.

Plain Word styles throughout — no layout tables, no text boxes — so it parses
cleanly in applicant tracking systems and stays editable for the recipient.

Usage:
    python3 make_docx.py
"""

import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from resume_content import (
    CONTACT_RIGHT,
    EDUCATION,
    EMAIL,
    EXPERIENCE,
    FOOTER,
    GITHUB,
    LINKEDIN,
    LOCATION,
    NAME,
    PROJECTS,
    ROLE,
    SKILLS,
    SUMMARY,
)

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Henok-Mesfin-Resume.docx")

BODY = "Calibri"
DISPLAY = "Georgia"
INK = RGBColor(0x1A, 0x1A, 0x17)
SOFT = RGBColor(0x33, 0x32, 0x2C)
MUTED = RGBColor(0x5C, 0x57, 0x4E)
ACCENT = RGBColor(0x8A, 0x5A, 0x16)
RULE = "D8D2C6"

PAGE_W = Inches(7.2)  # usable width: Letter - 0.65in margins


# ------------------------------------------------------------------ helpers
def style_run(run, size=9, bold=False, italic=False, color=SOFT, font=BODY):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    # make the font stick for non-Latin runs too
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return run


def para(doc, space_before=0, space_after=2, align=None, tab_right=True, line_spacing=None):
    p = doc.add_paragraph()
    if line_spacing is not None:
        p.paragraph_format.line_spacing = Pt(line_spacing)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    if tab_right:
        p.paragraph_format.tab_stops.add_tab_stop(PAGE_W, WD_TAB_ALIGNMENT.RIGHT)
    return p


def bottom_rule(p, color=RULE, size=6):
    """Thin horizontal rule under a heading paragraph."""
    ppr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    ppr.append(borders)


def add_hyperlink(paragraph, url, text, size=9):
    """Real Word hyperlink (python-docx has no built-in API for this)."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), BODY)
    rfonts.set(qn("w:hAnsi"), BODY)
    rPr.append(rfonts)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)

    col = OxmlElement("w:color")
    col.set(qn("w:val"), "8A5A16")
    rPr.append(col)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)

    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    link.append(new_run)
    paragraph._p.append(link)


def heading(doc, text):
    p = para(doc, space_before=7, space_after=2)
    style_run(p.add_run(text), size=9.5, bold=True, color=ACCENT)
    bottom_rule(p)
    return p


def rich(doc, parts, size=9, space_after=2, indent=None):
    """parts = [(text, bold, italic, color)]"""
    p = para(doc, space_after=space_after, line_spacing=11.5)
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    for text, bold, italic, color in parts:
        style_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p


# ------------------------------------------------------------------ sections
def build_header(doc):
    p = para(doc, space_after=0)
    style_run(p.add_run(NAME), size=19, bold=True, color=INK, font=DISPLAY)

    p = para(doc, space_after=1)
    style_run(p.add_run(ROLE), size=9, bold=True, color=ACCENT)

    p = para(doc, space_after=1)
    style_run(p.add_run("Email  "), size=9, color=MUTED)
    add_hyperlink(p, f"mailto:{EMAIL}", EMAIL)
    style_run(p.add_run("   ·   LinkedIn  "), size=9, color=MUTED)
    add_hyperlink(p, LINKEDIN, "linkedin.com/in/henok-mesfin-9341353b1")
    style_run(p.add_run("   ·   GitHub  "), size=9, color=MUTED)
    add_hyperlink(p, GITHUB, "github.com/henok-ne")
    style_run(p.add_run(f"   ·   {LOCATION}   ·   Open to remote contract work"), size=9, color=MUTED)
    return doc


def build_profile(doc):
    heading(doc, "PROFILE")
    rich(doc, [(SUMMARY, False, False, SOFT)], size=9, space_after=0)


def build_skills(doc):
    heading(doc, "TECHNICAL SKILLS")
    for label, value in SKILLS:
        rich(
            doc,
            [(f"{label}:  ", True, False, INK), (value, False, False, SOFT)],
            size=9,
            space_after=1,
        )


def build_experience(doc):
    heading(doc, "TECHNICAL EXPERIENCE")
    for job in EXPERIENCE:
        p = para(doc, space_before=4, space_after=0)
        style_run(p.add_run(job["title"]), size=9.5, bold=True, color=INK)
        style_run(p.add_run(f"\t{job['period']}"), size=9, color=MUTED)

        sub = "  ·  ".join(x for x in [job.get("context"), job.get("sub")] if x)
        if sub:
            p = para(doc, space_after=1)
            style_run(p.add_run(sub), size=9, italic=True, color=MUTED)

        for b in job["bullets"]:
            p = para(doc, space_after=1.5, tab_right=False, line_spacing=11.2)
            p.paragraph_format.left_indent = Inches(0.16)
            p.paragraph_format.first_line_indent = Inches(-0.16)
            style_run(p.add_run("–  " + b), size=9, color=SOFT)

        p = para(doc, space_after=0.5, tab_right=False)
        p.paragraph_format.left_indent = Inches(0.16)
        style_run(p.add_run(job["stack"]), size=8.5, italic=True, color=MUTED)


def build_projects(doc):
    heading(doc, "SELECTED PROJECTS")
    for title, desc, stack in PROJECTS:
        p = para(doc, space_after=2, tab_right=False, line_spacing=11.2)
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        style_run(p.add_run("–  "), size=9, color=SOFT)
        style_run(p.add_run(title), size=9, bold=True, color=INK)
        style_run(p.add_run(" — " + desc + " "), size=9, color=SOFT)
        style_run(p.add_run(stack), size=8.5, italic=True, color=MUTED)


def build_education(doc):
    heading(doc, "EDUCATION")
    for item in EDUCATION:
        p = para(doc, space_before=4, space_after=0)
        style_run(p.add_run(item["school"]), size=9.5, bold=True, color=INK)
        style_run(p.add_run(f"\t{item['right']}"), size=9, color=MUTED)
        rich(doc, [(item["detail"], False, False, SOFT)], size=9, space_after=0)


def build_footer(doc):
    p = para(doc, space_before=5, space_after=0)
    style_run(p.add_run(FOOTER), size=8, color=MUTED)


def build():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.45)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = BODY
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(2)

    build_header(doc)
    build_profile(doc)
    build_skills(doc)
    build_experience(doc)
    build_projects(doc)
    build_education(doc)
    build_footer(doc)

    doc.core_properties.title = "Henok Mesfin — Resume"
    doc.core_properties.author = "Henok Mesfin"
    doc.core_properties.subject = "College Software Engineering Expert — AfterQuery"

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
